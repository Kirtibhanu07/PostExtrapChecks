"""
QC Pipeline — Nielsen Branded Streamlit UI
Run: streamlit run app.py
"""

import re
import io
import time
import streamlit as st
import pandas as pd

from qc_check import (
    build_rate_check,
    build_exposure_check,
    build_sample,
    build_extrap_check,
    SHEETS,
)
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Page config ───────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="QC Pipeline · Nielsen",
    page_icon="🟣",
    layout="wide",
    initial_sidebar_state="collapsed",
)


def _html(content: str) -> None:
    flat = re.sub(r"\s*\n\s*", " ", content).strip()
    st.markdown(flat, unsafe_allow_html=True)


# ── CSS ───────────────────────────────────────────────────────────────────────
_html("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:ital,wght@0,300;0,400;0,500;0,600;0,700;0,800;1,400&display=swap');
@import url('https://fonts.googleapis.com/css2?family=Fraunces:ital,opsz,wght@0,9..144,400;0,9..144,500;0,9..144,600;0,9..144,700;1,9..144,500&display=swap');

/* ═══ NIELSEN TOKENS ═══ */
:root {
    --n-violet:       #6E37FA;
    --n-violet-dark:  #5324D9;
    --n-violet-wash:  #F3EFFE;
    --n-midnight:     #002041;
    --n-midnight-2:   #0B2A4D;
    --n-paper:        #F7F6F3;
    --n-panel:        #FFFFFF;
    --n-ink:          #201C2C;
    --n-ink-soft:     #6B6779;
    --n-ink-faint:    #A6A2B3;
    --n-line:         #E9E5F2;
    --n-line-soft:    #F1EEF8;
    --n-good:         #1F9D6C;
    --n-good-wash:    #EFFAF4;
    --n-fail:         #D6455A;
    --n-fail-wash:    #FDF1F2;
    --n-warn:         #C9821E;
    --n-warn-wash:    #FDF6EA;
    --n-maxw:         1280px; /* Expanded for less crampness */
}

/* ═══ RESET ═══ */
*, *::before, *::after { box-sizing: border-box; margin: 0; padding: 0; }

html, body,
[data-testid="stAppViewContainer"],
[data-testid="stAppViewContainer"] > section.main,
.block-container {
    background: var(--n-paper) !important;
    font-family: 'Inter', -apple-system, sans-serif !important;
    color: var(--n-ink) !important;
    max-width: 100% !important;
    padding: 0 !important;
}

[data-testid="stHeader"],
[data-testid="stToolbar"],
[data-testid="stDecoration"],
footer, #MainMenu { display: none !important; }

[data-testid="stHorizontalBlock"] { gap: 24px !important; align-items: stretch !important; }
[data-testid="column"] { padding: 0 !important; }
[data-testid="stVerticalBlock"] { gap: 0 !important; }

[data-testid="stAppViewContainer"] > .main .block-container {
    padding: 0 0 80px !important;
}

.n-shell { max-width: var(--n-maxw); margin: 0 auto; padding: 0 48px; }

/* ═══ TOPBAR ═══ */
.n-topbar-outer {
    background: linear-gradient(100deg, var(--n-midnight) 0%, var(--n-midnight-2) 100%);
    border-bottom: 3px solid var(--n-violet);
    margin-bottom: 56px;
    box-shadow: 0 4px 16px rgba(0,32,65,0.08);
}

.n-topbar {
    max-width: var(--n-maxw);
    margin: 0 auto;
    height: 84px;
    display: flex;
    align-items: center;
    justify-content: space-between;
    padding: 0 48px;
}

.n-topbar-left { display: flex; align-items: center; gap: 18px; }

.n-mark {
    width: 38px; height: 38px;
    border-radius: 10px;
    background: var(--n-violet);
    display: flex; align-items: center; justify-content: center;
    font-family: 'Fraunces', serif;
    font-style: italic;
    font-weight: 700;
    font-size: 19px;
    color: #FFFFFF;
    flex-shrink: 0;
    box-shadow: 0 2px 8px rgba(110,55,250,0.4);
}

.n-topbar-text { display: flex; flex-direction: column; gap: 3px; }

.n-wordmark {
    font-family: 'Fraunces', serif;
    font-size: 19px;
    font-weight: 600;
    font-style: italic;
    color: #FFFFFF;
    letter-spacing: -0.1px;
    line-height: 1;
}

.n-app-name { font-size: 13px; font-weight: 500; color: rgba(255,255,255,0.65); letter-spacing: 0.01em; }

.n-version {
    font-size: 12px;
    font-weight: 600;
    color: #FFFFFF;
    font-family: 'Inter', monospace;
    background: rgba(110,55,250,0.35);
    border: 1px solid rgba(110,55,250,0.6);
    padding: 6px 14px;
    border-radius: 20px;
}

/* ═══ SECTION HEADER ═══ */
.n-step-head {
    display: flex;
    align-items: baseline;
    gap: 16px;
    margin-bottom: 20px;
}

.n-step-num {
    font-family: 'Fraunces', serif;
    font-style: italic;
    font-size: 14px;
    font-weight: 600;
    color: var(--n-violet);
    background: var(--n-violet-wash);
    width: 28px; height: 28px;
    border-radius: 50%;
    display: flex; align-items: center; justify-content: center;
    flex-shrink: 0;
    transform: translateY(1px);
}

.n-step-title {
    font-family: 'Fraunces', serif;
    font-size: 20px;
    font-weight: 500;
    color: var(--n-ink);
}

.n-step-sub { font-size: 13.5px; color: var(--n-ink-faint); margin-left: 44px; margin-top: -14px; margin-bottom: 28px; }

.n-section { margin-bottom: 56px; }

/* ═══ UPLOAD GRID ═══ */
.n-upload-card {
    background: var(--n-panel);
    border: 1px solid var(--n-line);
    border-radius: 16px;
    padding: 22px;
    height: 100%;
    display: flex;
    flex-direction: column;
    transition: border-color 0.2s, box-shadow 0.2s;
    box-shadow: 0 4px 16px rgba(32,28,44,0.02);
}

.n-upload-card.is-loaded { 
    border-color: rgba(31,157,108,0.4); 
    box-shadow: 0 4px 20px rgba(31,157,108,0.08); 
}

.n-upload-card-head {
    display: flex;
    align-items: center;
    justify-content: space-between;
    margin-bottom: 6px;
}

.n-upload-name { font-size: 14px; font-weight: 700; color: var(--n-ink); }
.n-upload-hint { font-size: 11.5px; color: var(--n-ink-faint); font-family: 'Inter', monospace; margin-bottom: 14px; }

.n-upload-dot { width: 8px; height: 8px; border-radius: 50%; background: #E4E1EC; flex-shrink: 0; }
.n-upload-dot.loaded { background: var(--n-good); box-shadow: 0 0 0 4px rgba(31,157,108,0.15); }

.n-upload-status {
    margin-top: 14px;
    font-size: 12px;
    font-weight: 600;
    color: var(--n-ink-faint);
}
.n-upload-status.loaded { color: var(--n-good); }

/* Compact uploader */
.n-upload-card [data-testid="stFileUploader"] { background: transparent !important; margin: 0 !important; padding: 0 !important; }
.n-upload-card [data-testid="stFileUploader"] section { padding: 0 !important; }
.n-upload-card [data-testid="stFileUploaderDropzoneInstructions"] { display: none !important; }

[data-testid="stFileUploaderDropzone"],
[data-testid="stFileUploader"] section {
    background: #FBFAF8 !important;
    border: 1.5px dashed #DDD7EE !important;
    border-radius: 10px !important;
    padding: 24px 16px !important;
    min-height: unset !important;
    display: flex !important;
    align-items: center !important;
    justify-content: center !important;
    transition: border-color 0.2s, background 0.2s !important;
    cursor: pointer !important;
    position: relative !important;
}

[data-testid="stFileUploaderDropzone"]:hover,
[data-testid="stFileUploader"] section:hover {
    border-color: var(--n-violet) !important;
    background: var(--n-violet-wash) !important;
}

[data-testid="stFileUploaderFileName"] { color: var(--n-ink) !important; font-size: 11.5px !important; font-weight: 500 !important; }

.n-upload-card [data-testid="stFileUploader"] button,
.n-upload-card [data-testid="stFileUploaderDropzone"] button,
.n-upload-card [data-testid="baseButton-secondary"],
.n-upload-card [data-testid="stBaseButton-secondary"] {
    visibility: hidden !important;
    position: absolute !important;
    width: 1px !important;
    height: 1px !important;
    pointer-events: none !important;
}

[data-testid="stFileUploaderDropzone"]::after,
[data-testid="stFileUploader"] section::after {
    content: 'Click or drop .xlsx here';
    font-size: 12px;
    font-weight: 500;
    color: var(--n-ink-faint);
    pointer-events: none;
}

[data-testid="stFileUploaderDropzone"]:hover::after,
[data-testid="stFileUploader"] section:hover::after {
    color: var(--n-violet-dark);
}

[data-testid="stFileUploaderFile"] ~ *::after,
[data-testid="stFileUploaderDropzone"]:has([data-testid="stFileUploaderFile"])::after {
    content: '' !important;
}

/* ═══ ACTION ROW ═══ */
.n-action-row {
    display: flex;
    align-items: center;
    gap: 36px;
    background: var(--n-panel);
    border: 1px solid var(--n-line);
    border-radius: 16px;
    padding: 26px 36px;
    margin-top: 24px;
    box-shadow: 0 4px 16px rgba(32,28,44,0.02);
}

.n-action-progress { flex: 1; }

.n-progress-track { height: 8px; background: var(--n-line-soft); border-radius: 6px; overflow: hidden; margin-bottom: 12px; }
.n-progress-fill { height: 100%; background: linear-gradient(90deg, var(--n-violet), var(--n-violet-dark)); border-radius: 6px; transition: width 0.4s ease; }
.n-progress-label { font-size: 13.5px; color: var(--n-ink-faint); font-weight: 500; }
.n-progress-label span { color: var(--n-ink); font-weight: 700; }

.n-action-row [data-testid="stButton"] { min-width: 280px; }

[data-testid="stButton"] > button {
    width: 100% !important;
    height: 54px !important;
    background: linear-gradient(135deg, var(--n-violet) 0%, var(--n-violet-dark) 100%) !important;
    color: #FFFFFF !important;
    font-family: 'Inter', sans-serif !important;
    font-size: 14.5px !important;
    font-weight: 700 !important;
    letter-spacing: 0.03em !important;
    border: none !important;
    border-radius: 12px !important;
    cursor: pointer !important;
    transition: transform 0.15s ease, box-shadow 0.2s !important;
    box-shadow: 0 4px 14px rgba(110,55,250,0.35), inset 0 1px 0 rgba(255,255,255,0.18) !important;
}

[data-testid="stButton"] > button:hover:not(:disabled) {
    transform: translateY(-2px);
    box-shadow: 0 8px 22px rgba(110,55,250,0.45), inset 0 1px 0 rgba(255,255,255,0.2) !important;
}

[data-testid="stButton"] > button:disabled {
    background: #F1EFF7 !important;
    color: #BBB6CC !important;
    box-shadow: none !important;
    cursor: not-allowed !important;
    border: 1px solid var(--n-line) !important;
}

/* ═══ PIPELINE STEPPER ═══ */
.n-pipeline-bar { 
    background: var(--n-panel); 
    border: 1px solid var(--n-line); 
    border-radius: 16px; 
    padding: 36px 44px; 
    box-shadow: 0 4px 20px rgba(32,28,44,0.03); 
}

.n-stages { display: grid; grid-template-columns: repeat(4, 1fr); gap: 24px; position: relative; }

.n-stages::before {
    content: '';
    position: absolute;
    top: 20px;
    left: calc(12.5% + 20px);
    right: calc(12.5% + 20px);
    height: 2px;
    background: var(--n-line-soft);
    z-index: 0;
}

.n-stage-col { display: flex; flex-direction: column; align-items: center; gap: 14px; position: relative; z-index: 1; }

.n-stage-icon {
    width: 42px; height: 42px;
    border-radius: 50%;
    border: 2px solid #E4E1EC;
    background: #FFFFFF;
    display: flex; align-items: center; justify-content: center;
    font-size: 16px;
    font-weight: 700;
    color: #C9C4DB;
    transition: all 0.3s ease;
}

.n-stage-col.done   .n-stage-icon { border-color: var(--n-good); background: var(--n-good); color: #FFFFFF; box-shadow: 0 4px 12px rgba(31,157,108,0.2); }
.n-stage-col.active .n-stage-icon { border-color: var(--n-violet); background: var(--n-violet); color: #FFFFFF; box-shadow: 0 0 0 5px rgba(110,55,250,0.14); }
.n-stage-col.failed .n-stage-icon { border-color: var(--n-fail); background: var(--n-fail-wash); color: var(--n-fail); }
.n-stage-col.skipped .n-stage-icon { border-color: var(--n-warn); background: var(--n-warn-wash); color: var(--n-warn); }

.n-stage-label { font-size: 14px; font-weight: 600; color: var(--n-ink-faint); text-align: center; transition: color 0.3s; }
.n-stage-col.done   .n-stage-label { color: var(--n-good); }
.n-stage-col.active .n-stage-label { color: var(--n-violet); }
.n-stage-col.failed .n-stage-label { color: var(--n-fail); }
.n-stage-col.skipped .n-stage-label { color: var(--n-warn); }

.n-stage-sub { font-size: 12px; color: #C9C4DB; text-align: center; line-height: 1.5; }
.n-stage-col.done .n-stage-sub   { color: var(--n-ink-faint); }
.n-stage-col.active .n-stage-sub { color: var(--n-violet); opacity: 0.8; }

/* ═══ LOG + RESULTS ROW ═══ */
.n-row-header {
    display: flex;
    align-items: center;
    justify-content: space-between;
    margin-bottom: 24px;
}

.n-toggle-wrap [data-testid="stWidgetLabel"] p { font-size: 13px !important; font-weight: 600 !important; color: var(--n-ink-soft) !important; }

.n-log-card { 
    background: var(--n-midnight); 
    border-radius: 16px; 
    overflow: hidden; 
    box-shadow: 0 8px 32px rgba(0,32,65,0.08); 
}

.n-log-header {
    padding: 20px 32px;
    border-bottom: 1px solid rgba(255,255,255,0.08);
    display: flex;
    align-items: center;
    justify-content: space-between;
    background: rgba(255,255,255,0.03);
}

.n-log-header-title { font-size: 12.5px; font-weight: 600; color: rgba(255,255,255,0.6); letter-spacing: 0.1em; text-transform: uppercase; }

.n-log-dot { width: 8px; height: 8px; border-radius: 50%; background: rgba(255,255,255,0.18); }
.n-log-dot.active { background: var(--n-violet); box-shadow: 0 0 8px rgba(110,55,250,0.7); }
.n-log-dot.done   { background: var(--n-good); }

.n-log-body {
    padding: 28px 32px;
    font-family: 'SF Mono', 'Fira Code', 'Consolas', monospace;
    font-size: 13.5px;
    line-height: 2;
    min-height: 240px;
    max-height: 400px;
    overflow-y: auto;
}

.n-log-body::-webkit-scrollbar { width: 6px; }
.n-log-body::-webkit-scrollbar-track { background: rgba(255,255,255,0.02); }
.n-log-body::-webkit-scrollbar-thumb { background: rgba(255,255,255,0.15); border-radius: 3px; }

.ll { display: flex; gap: 14px; }
.ll-ts { color: rgba(255,255,255,0.25); flex-shrink: 0; }
.ll-ok   { color: #8FB8FF; }
.ll-good { color: #5CDDA0; }
.ll-fail { color: #FF7A8A; }
.ll-warn { color: #F0B84E; }
.ll-dim  { color: rgba(255,255,255,0.25); }
.ll-head { color: #B296FF; font-weight: 600; }
.ll-stage { color: rgba(255,255,255,0.4); }
.ll-div { border: none; border-top: 1px solid rgba(255,255,255,0.08); margin: 8px 0; }

.n-empty { display: flex; flex-direction: column; align-items: center; justify-content: center; gap: 12px; padding: 60px 48px; text-align: center; }
.n-empty-icon { font-size: 28px; opacity: 0.5; color: var(--n-violet); }
.n-empty-title { font-family: 'Fraunces', serif; font-size: 16px; font-weight: 500; font-style: italic; color: rgba(255,255,255,0.8); }
.n-empty-sub { font-size: 13px; color: rgba(255,255,255,0.4); max-width: 300px; line-height: 1.6; }

.n-log-compact {
    background: var(--n-panel);
    border: 1px solid var(--n-line);
    border-radius: 16px;
    padding: 26px 32px;
    display: flex;
    align-items: center;
    gap: 18px;
    box-shadow: 0 4px 16px rgba(32,28,44,0.02);
}

.n-log-compact-dot { width: 10px; height: 10px; border-radius: 50%; background: #E4E1EC; flex-shrink: 0; }
.n-log-compact-dot.active { background: var(--n-violet); box-shadow: 0 0 0 5px rgba(110,55,250,0.14); }
.n-log-compact-dot.done   { background: var(--n-good); box-shadow: 0 0 0 5px rgba(31,157,108,0.14); }
.n-log-compact-dot.fail   { background: var(--n-fail); box-shadow: 0 0 0 5px rgba(214,69,90,0.14); }

.n-log-compact-text { font-size: 14.5px; font-weight: 600; color: var(--n-ink); }
.n-log-compact-sub { font-size: 12.5px; color: var(--n-ink-faint); margin-top: 3px; }

/* ═══ RESULTS ═══ */
.n-results-head {
    display: flex;
    align-items: center;
    justify-content: space-between;
    margin-bottom: 24px;
}

.n-result-grid { display: grid; grid-template-columns: repeat(4, 1fr); gap: 20px; }

.n-result-card { border-radius: 16px; padding: 24px 28px; border: 1px solid var(--n-line-soft); background: #FCFBFA; box-shadow: 0 4px 16px rgba(32,28,44,0.02); }
.n-result-card.ok   { border-color: rgba(31,157,108,0.3); background: var(--n-good-wash); }
.n-result-card.fail { border-color: rgba(214,69,90,0.3);  background: var(--n-fail-wash); }
.n-result-card.skip { border-color: rgba(201,130,30,0.3); background: var(--n-warn-wash); }

.n-rc-name { font-size: 11px; font-weight: 700; letter-spacing: 0.1em; text-transform: uppercase; color: var(--n-ink-faint); margin-bottom: 12px; }
.n-rc-stat { font-size: 26px; font-weight: 800; line-height: 1; margin-bottom: 10px; }
.n-result-card.ok   .n-rc-stat { color: var(--n-good); }
.n-result-card.fail .n-rc-stat { color: var(--n-fail); }
.n-result-card.skip .n-rc-stat { color: var(--n-warn); }
.n-rc-sub { font-size: 12px; color: var(--n-ink-faint); line-height: 1.5; }

[data-testid="stDownloadButton"] > button {
    height: 52px !important;
    background: var(--n-midnight) !important;
    color: #FFFFFF !important;
    font-family: 'Inter', sans-serif !important;
    font-size: 14px !important;
    font-weight: 600 !important;
    border: none !important;
    border-radius: 10px !important;
    padding: 0 32px !important;
    letter-spacing: 0.02em !important;
    cursor: pointer !important;
    transition: background 0.2s, box-shadow 0.2s !important;
    white-space: nowrap !important;
    box-shadow: 0 4px 12px rgba(0,32,65,0.15) !important;
}

[data-testid="stDownloadButton"] > button:hover { 
    background: var(--n-midnight-2) !important; 
    box-shadow: 0 6px 16px rgba(0,32,65,0.2) !important;
}

[data-testid="stSpinner"] > div { color: var(--n-violet) !important; }
.uploadedFileName { color: var(--n-ink) !important; font-size: 12px !important; }
</style>
""")

# ── Session state ─────────────────────────────────────────────────────────────
if "stage_states" not in st.session_state:
    st.session_state.stage_states = {1:"idle", 2:"idle", 3:"idle", 4:"idle"}
if "log_lines" not in st.session_state:
    st.session_state.log_lines = []
if "result_stats" not in st.session_state:
    st.session_state.result_stats = None
if "output_bytes" not in st.session_state:
    st.session_state.output_bytes = None
if "show_logs" not in st.session_state:
    st.session_state.show_logs = True

# ── Helpers ───────────────────────────────────────────────────────────────────
def ts():
    return time.strftime("%H:%M:%S")

def add_log(text, kind="ok"):
    st.session_state.log_lines.append((ts(), text, kind))

def set_stage(n, state):
    st.session_state.stage_states[n] = state

# ── Excel writer ──────────────────────────────────────────────────────────────
def write_output_styled(sheets: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    HDR_FILL  = PatternFill("solid", fgColor="1A1A1A")
    HDR_FONT  = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    HDR_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
    CELL_FONT = Font(name="Calibri", size=10)
    ODD_FILL  = PatternFill("solid", fgColor="F9F9F9")
    EVN_FILL  = PatternFill("solid", fgColor="FFFFFF")
    T_FILL    = PatternFill("solid", fgColor="F0FBF6"); T_FONT = Font(name="Calibri", size=10, color="22A06B", bold=True)
    F_fill    = PatternFill("solid", fgColor="FFF5F5"); F_FONT = Font(name="Calibri", size=10, color="CC0000", bold=True)
    thin      = Side(style="thin", color="EEEEEE")
    BORDER    = Border(bottom=thin)

    for name, df in sheets.items():
        ws = wb.create_sheet(title=name)
        ws.sheet_properties.tabColor = "FFFF00"

        for ci, col in enumerate(df.columns, 1):
            c = ws.cell(row=1, column=ci, value=str(col))
            c.fill=HDR_FILL; c.font=HDR_FONT; c.alignment=HDR_ALIGN; c.border=BORDER

        for ri, row in enumerate(df.itertuples(index=False), 2):
            fill = ODD_FILL if ri % 2 == 0 else EVN_FILL
            for ci, val in enumerate(row, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                cname = df.columns[ci-1]
                if cname == "Check":
                    if val is True:
                        c.fill=T_FILL; c.font=T_FONT; c.value="TRUE"
                    elif val is False:
                        c.fill=F_fill; c.font=F_FONT; c.value="FALSE"
                    else:
                        c.fill=fill; c.font=CELL_FONT
                else:
                    c.fill=fill; c.font=CELL_FONT

        for col in ws.columns:
            ltr = get_column_letter(col[0].column)
            mx = max((len(str(c.value)) for c in col if c.value), default=8)
            ws.column_dimensions[ltr].width = min(max(mx + 2, 8), 36)

        ws.freeze_panes = "A2"
        if df.shape[1] > 0:
            ws.auto_filter.ref = f"A1:{get_column_letter(df.shape[1])}1"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

# ── Pipeline runner ───────────────────────────────────────────────────────────
def run_pipeline(adapt_f, bsr_f, yt_f, matex_f, sample_f):
    st.session_state.log_lines = []
    st.session_state.result_stats = None
    st.session_state.output_bytes = None
    for n in [1,2,3,4]: set_stage(n, "idle")

    add_log("QC CHECK AUTOMATION PIPELINE", "head")
    add_log(f"Started at {time.strftime('%Y-%m-%d %H:%M:%S')}", "dim")
    add_log("", "dim")
    add_log("── LOAD ──────────────────────────────────────────────", "stage")

    load_ok = True
    frames  = {}

    for key, uf, sheet, hdr, label in [
        ("adapt",  adapt_f,  SHEETS["adapt"],  0, "Adapt"),
        ("bsr",    bsr_f,    SHEETS["bsr"],    5, "BSR"),
        ("yt",     yt_f,     SHEETS["yt"],     0, "YouTube MM"),
        ("matex",  matex_f,  SHEETS["matex"],  0, "Matex"),
        ("sample", sample_f, SHEETS["sample"], 0, "Sample"),
    ]:
        try:
            buf = io.BytesIO(uf.read())
            df  = pd.read_excel(buf, sheet_name=sheet, header=hdr, engine="openpyxl")
            df.columns = [str(c).strip() for c in df.columns]
            frames[key] = df
            add_log(f"  [✓]  {label:<14}  {len(df):,} rows · {len(df.columns)} cols", "good")
        except Exception as e:
            add_log(f"  [✗]  {label:<14}  {e}", "fail")
            load_ok = False

    if not load_ok:
        add_log("", "dim"); add_log("  PIPELINE ABORTED — fix load errors above", "fail")
        return

    add_log("", "dim")
    add_log("── STAGES ────────────────────────────────────────────", "stage")

    output_sheets  = {}
    sample_summary = None
    stats = {}

    set_stage(1, "active")
    try:
        df = build_rate_check(frames["adapt"], frames["bsr"], frames["yt"])
        output_sheets["RateCheck"] = df
        stats[1] = {"rows": len(df), "ok": True}
        set_stage(1, "done")
        add_log(f"  [✓]  RateCheck        {len(df):,} rows", "good")
    except Exception as e:
        set_stage(1, "failed"); stats[1] = {"ok": False, "error": str(e)}
        add_log(f"  [✗]  RateCheck        {e}", "fail")

    set_stage(2, "active")
    try:
        df = build_exposure_check(frames["sample"], frames["matex"])
        output_sheets["ExposureCheck"] = df
        stats[2] = {"rows": len(df), "ok": True}
        set_stage(2, "done")
        add_log(f"  [✓]  ExposureCheck    {len(df):,} rows", "good")
    except Exception as e:
        set_stage(2, "failed"); stats[2] = {"ok": False, "error": str(e)}
        add_log(f"  [✗]  ExposureCheck    {e}", "fail")

    set_stage(3, "active")
    try:
        ss, sample_summary = build_sample(frames["sample"])
        output_sheets["Sample"] = ss
        stats[3] = {"rows": len(ss), "matchdays": len(sample_summary), "ok": True}
        set_stage(3, "done")
        add_log(f"  [✓]  Sample           {len(ss):,} rows · {len(sample_summary)} matchdays", "good")
    except Exception as e:
        set_stage(3, "failed"); stats[3] = {"ok": False, "error": str(e)}
        add_log(f"  [✗]  Sample           {e}", "fail")

    set_stage(4, "active")
    try:
        if sample_summary is None:
            raise RuntimeError("Stage 3 must succeed before Stage 4 can run")
        df = build_extrap_check(frames["adapt"], sample_summary)
        output_sheets["ExtrapCheck"] = df
        stats[4] = {"rows": len(df), "ok": True}
        set_stage(4, "done")
        add_log(f"  [✓]  ExtrapCheck      {len(df):,} rows", "good")
    except RuntimeError as e:
        set_stage(4, "skipped"); stats[4] = {"ok": False, "skipped": True, "error": str(e)}
        add_log(f"  [⚠]  ExtrapCheck      SKIPPED — {e}", "warn")
    except Exception as e:
        set_stage(4, "failed"); stats[4] = {"ok": False, "error": str(e)}
        add_log(f"  [✗]  ExtrapCheck      {e}", "fail")

    add_log("", "dim")
    add_log("── OUTPUT ────────────────────────────────────────────", "stage")
    done = sum(1 for s in stats.values() if s.get("ok"))

    if output_sheets:
        try:
            st.session_state.output_bytes = write_output_styled(output_sheets)
            add_log(f"  [✓]  {done}/4 sheets written → QC_Output.xlsx", "good")
        except Exception as e:
            add_log(f"  [✗]  Write failed: {e}", "fail")

    add_log("", "dim")
    if done == 4:
        add_log("  ✓  COMPLETE — all 4 sheets written", "good")
    elif done > 0:
        add_log(f"  ⚠  PARTIAL — {done}/4 sheets written", "warn")
    else:
        add_log("  ✗  FAILED — no output written", "fail")

    st.session_state.result_stats = stats


# ══════════════════════════════════════════════════════════════════════════════
# RENDER
# ══════════════════════════════════════════════════════════════════════════════

_html("""
<div class="n-topbar-outer">
  <div class="n-topbar">
    <div class="n-topbar-left">
      <div class="n-mark">N</div>
      <div class="n-topbar-text">
        <span class="n-wordmark">Nielsen</span>
        <span class="n-app-name">QC Automation Pipeline</span>
      </div>
    </div>
    <span class="n-version">v7.0 · Media Measurement</span>
  </div>
</div>
""")

_html('<div class="n-shell">')

# ═══════════════════════════════════════════════════════════════════════════
# STEP 1
# ═══════════════════════════════════════════════════════════════════════════
_html('<div class="n-section">')
_html("""
<div class="n-step-head">
  <div class="n-step-num">1</div>
  <div class="n-step-title">Source Files</div>
</div>
<div class="n-step-sub">Upload all five workbooks to unlock the pipeline run</div>
""")

FILE_DEFS = [
    ("adapt",  "Adapt",      "programs"),
    ("bsr",    "BSR",        "Database · row 6"),
    ("yt",     "YouTube MM", "MM"),
    ("matex",  "Matex",      "exposures"),
    ("sample", "Sample",     "programs"),
]

uploaded = {}
loaded_count = 0
upload_cols = st.columns(5, gap="medium") # Increased gap for premium feel

for col, (key, label, hint) in zip(upload_cols, FILE_DEFS):
    with col:
        uf = st.session_state.get(f"upload_{key}")
        loaded_now = uf is not None
        _html(f'<div class="n-upload-card{" is-loaded" if loaded_now else ""}">')
        _html(f"""
        <div class="n-upload-card-head">
          <span class="n-upload-name">{label}</span>
          <div class="n-upload-dot{' loaded' if loaded_now else ''}"></div>
        </div>
        <div class="n-upload-hint">sheet: {hint}</div>
        """)
        uf = st.file_uploader(
            label, type=["xlsx"], key=f"upload_{key}",
            help=f"Sheet: {hint}", label_visibility="collapsed",
        )
        uploaded[key] = uf
        if uf:
            loaded_count += 1
            status_txt = f"✓ {uf.size/1024:.0f} KB"
        else:
            status_txt = "Awaiting upload"
        _html(f'<div class="n-upload-status{" loaded" if uf else ""}">{status_txt}</div>')
        _html('</div>')

pct = int(loaded_count / 5 * 100)
all_ready = loaded_count == 5
btn_label = "Run QC Pipeline →" if all_ready else f"Waiting for {5 - loaded_count} more file{'s' if 5-loaded_count != 1 else ''}"

_html('<div class="n-action-row">')
action_progress_col, action_btn_col = st.columns([2.5, 1], gap="large")
with action_progress_col:
    _html(f"""
    <div class="n-action-progress">
      <div class="n-progress-track">
        <div class="n-progress-fill" style="width:{pct}%"></div>
      </div>
      <div class="n-progress-label"><span>{loaded_count}</span> of 5 files ready</div>
    </div>
    """)
with action_btn_col:
    if st.button(btn_label, disabled=not all_ready, use_container_width=True):
        with st.spinner("Running…"):
            run_pipeline(
                uploaded["adapt"], uploaded["bsr"], uploaded["yt"],
                uploaded["matex"], uploaded["sample"]
            )
        st.rerun()
_html('</div>')
_html('</div>')

# ═══════════════════════════════════════════════════════════════════════════
# STEP 2
# ═══════════════════════════════════════════════════════════════════════════
_html('<div class="n-section">')
_html("""
<div class="n-step-head">
  <div class="n-step-num">2</div>
  <div class="n-step-title">Pipeline Stages</div>
</div>
""")

STAGE_DEFS = [
    (1, "RateCheck",     "Adapt · BSR · YT"),
    (2, "ExposureCheck", "Sample vs Matex"),
    (3, "Sample",        "pID · matchdays"),
    (4, "ExtrapCheck",   "BT% / Msec%"),
]

ICON_MAP = {"idle":"—","done":"✓","active":"◎","failed":"✗","skipped":"⚠"}

cols_html = ""
for n, title, sub in STAGE_DEFS:
    state = st.session_state.stage_states[n]
    icon  = ICON_MAP.get(state, "—")

    if state == "done" and st.session_state.result_stats:
        s = st.session_state.result_stats.get(n, {})
        rows = s.get("rows")
        if rows:
            sub = f"{rows:,} rows"
        if n == 3 and s.get("matchdays"):
            sub += f" · {s['matchdays']} matchdays"

    cols_html += f"""
    <div class="n-stage-col {state}">
      <div class="n-stage-icon">{icon}</div>
      <div class="n-stage-label">{title}</div>
      <div class="n-stage-sub">{sub}</div>
    </div>"""

_html(f"""
<div class="n-pipeline-bar">
  <div class="n-stages">{cols_html}</div>
</div>
""")
_html('</div>')

# ═══════════════════════════════════════════════════════════════════════════
# STEP 3
# ═══════════════════════════════════════════════════════════════════════════
_html('<div class="n-section">')

done_count = sum(1 for s in (st.session_state.result_stats or {}).values() if s.get("ok"))
dot_class = "active" if st.session_state.log_lines and not st.session_state.result_stats else ("done" if st.session_state.result_stats else "")

head_col1, head_col2 = st.columns([3, 1], gap="medium")
with head_col1:
    _html("""
    <div class="n-step-head" style="margin-bottom:0;">
      <div class="n-step-num">3</div>
      <div class="n-step-title">Pipeline Log</div>
    </div>
    """)
with head_col2:
    _html('<div class="n-toggle-wrap" style="display:flex; justify-content:flex-end;">')
    st.session_state.show_logs = st.toggle(
        "Show logs", value=st.session_state.show_logs, key="show_logs_toggle",
    )
    _html('</div>')

_html('<div style="height:22px"></div>')

if st.session_state.show_logs:
    if not st.session_state.log_lines:
        log_content = """
        <div class="n-empty">
          <div class="n-empty-icon">○</div>
          <div class="n-empty-title">No pipeline run yet</div>
          <div class="n-empty-sub">Upload all 5 source files and press Run QC Pipeline</div>
        </div>"""
    else:
        rows_html = ""
        for stamp, text, kind in st.session_state.log_lines:
            if text == "":
                rows_html += '<div style="height:4px"></div>'
            elif text.startswith("──"):
                rows_html += f'<hr class="ll-div"><div class="ll"><span class="ll-ts">{stamp}</span><span class="ll-stage">{text}</span></div>'
            else:
                cls = {"ok":"ll-ok","good":"ll-good","fail":"ll-fail","warn":"ll-warn","dim":"ll-dim","head":"ll-head","stage":"ll-stage"}.get(kind,"ll-ok")
                rows_html += f'<div class="ll"><span class="ll-ts">{stamp}</span><span class="{cls}">{text}</span></div>'
        log_content = f'<div class="n-log-body">{rows_html}</div>'

    _html(f"""
    <div class="n-log-card">
      <div class="n-log-header">
        <span class="n-log-header-title">Live Output</span>
        <div class="n-log-dot {dot_class}"></div>
      </div>
      {log_content}
    </div>
    """)
else:
    if not st.session_state.log_lines:
        compact_dot, compact_text, compact_sub = "", "No pipeline run yet", "Toggle logs on for step-by-step detail"
    elif st.session_state.result_stats and done_count == 4:
        compact_dot, compact_text, compact_sub = "done", "Complete — all 4 sheets written", "Toggle logs on for step-by-step detail"
    elif st.session_state.result_stats and done_count > 0:
        compact_dot, compact_text, compact_sub = "fail", f"Partial — {done_count}/4 sheets written", "Toggle logs on to see what failed"
    elif st.session_state.result_stats:
        compact_dot, compact_text, compact_sub = "fail", "Failed — no output written", "Toggle logs on to see what failed"
    else:
        compact_dot, compact_text, compact_sub = "active", "Running…", "Toggle logs on for step-by-step detail"

    _html(f"""
    <div class="n-log-compact">
      <div class="n-log-compact-dot {compact_dot}"></div>
      <div>
        <div class="n-log-compact-text">{compact_text}</div>
        <div class="n-log-compact-sub">{compact_sub}</div>
      </div>
    </div>
    """)

_html('</div>')

# ═══════════════════════════════════════════════════════════════════════════
# STEP 4
# ═══════════════════════════════════════════════════════════════════════════
if st.session_state.result_stats:
    _html('<div class="n-section">')

    res_head_col1, res_head_col2 = st.columns([3, 1], gap="medium")
    with res_head_col1:
        _html("""
        <div class="n-step-head" style="margin-bottom:0;">
          <div class="n-step-num">4</div>
          <div class="n-step-title">Output Sheets</div>
        </div>
        """)
    with res_head_col2:
        if st.session_state.output_bytes:
            fname = f"QC_Output_{time.strftime('%Y%m%d_%H%M')}.xlsx"
            st.download_button(
                label="↓  Download",
                data=st.session_state.output_bytes,
                file_name=fname,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )

    _html('<div style="height:22px"></div>')

    stats = st.session_state.result_stats
    names = ["RateCheck","ExposureCheck","Sample","ExtrapCheck"]

    cards_html = ""
    for i, sname in enumerate(names, 1):
        s = stats.get(i, {})
        ok      = s.get("ok", False)
        skipped = s.get("skipped", False)
        cls     = "ok" if ok else ("skip" if skipped else "fail")
        icon    = "✓" if ok else ("⚠" if skipped else "✗")
        rows    = s.get("rows")
        if ok:
            sub = f"{rows:,} rows" if rows else "complete"
            if i == 3 and s.get("matchdays"): sub += f"<br>{s['matchdays']} matchdays"
        elif skipped:
            sub = "Stage 3 required"
        else:
            sub = "See log above"

        cards_html += f"""
        <div class="n-result-card {cls}">
          <div class="n-rc-name">{sname}</div>
          <div class="n-rc-stat">{icon}</div>
          <div class="n-rc-sub">{sub}</div>
        </div>"""

    _html(f'<div class="n-result-grid">{cards_html}</div>')
    _html('</div>')

_html('</div>')
