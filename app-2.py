"""
QC Pipeline — Nielsen Branded Streamlit UI
Run: streamlit run app.py
"""

import io
import time
import streamlit as st
import pandas as pd

from qc_check import (
    build_rate_check,
    build_exposure_check,
    build_sample,
    build_extrap_check,
    SHEETS,
)
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Page config ───────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="QC Pipeline · Nielsen",
    page_icon="🟣",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ── CSS ───────────────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:ital,wght@0,300;0,400;0,500;0,600;0,700;0,800;1,400&display=swap');
@import url('https://fonts.googleapis.com/css2?family=Fraunces:ital,opsz,wght@0,9..144,400;0,9..144,500;0,9..144,600;0,9..144,700;1,9..144,500&display=swap');

/* ═══ NIELSEN TOKENS (2021 brand refresh: violet/midnight, not red) ═══ */
:root {
    --n-violet:       #6E37FA;
    --n-violet-dark:  #5324D9;
    --n-violet-wash:  #F3EFFE;
    --n-midnight:     #002041;
    --n-midnight-2:   #0B2A4D;
    --n-paper:        #FAF9F6;
    --n-panel:        #FFFFFF;
    --n-ink:          #201C2C;
    --n-ink-soft:     #6B6779;
    --n-ink-faint:    #A6A2B3;
    --n-line:         #E9E5F2;
    --n-line-soft:    #F1EEF8;
    --n-good:         #1F9D6C;
    --n-good-wash:    #EFFAF4;
    --n-fail:         #D6455A;
    --n-fail-wash:    #FDF1F2;
    --n-warn:         #C9821E;
    --n-warn-wash:    #FDF6EA;
}

/* ═══ RESET ═══ */
*, *::before, *::after { box-sizing: border-box; margin: 0; padding: 0; }

html, body,
[data-testid="stAppViewContainer"],
[data-testid="stAppViewContainer"] > section.main,
.block-container {
    background: var(--n-paper) !important;
    background-image: radial-gradient(circle at 15% 0%, rgba(110,55,250,0.05), transparent 40%),
                       radial-gradient(circle at 100% 30%, rgba(0,32,65,0.035), transparent 45%) !important;
    font-family: 'Inter', -apple-system, sans-serif !important;
    color: var(--n-ink) !important;
    max-width: 100% !important;
    padding: 0 !important;
}

/* Kill all Streamlit default chrome */
[data-testid="stHeader"],
[data-testid="stToolbar"],
[data-testid="stDecoration"],
footer, #MainMenu { display: none !important; }

[data-testid="stHorizontalBlock"] { gap: 32px !important; align-items: stretch !important; }
[data-testid="column"] { padding: 0 !important; gap: 0 !important; }
[data-testid="stVerticalBlock"] { gap: 0 !important; }

/* Streamlit wraps page content in a centered block with side padding by default */
[data-testid="stAppViewContainer"] > .main .block-container {
    padding-top: 0 !important;
    padding-bottom: 40px !important;
    padding-left: 40px !important;
    padding-right: 40px !important;
}

/* ═══ TOPBAR ═══ */
.n-topbar {
    background: linear-gradient(100deg, var(--n-midnight) 0%, var(--n-midnight-2) 100%);
    height: 68px;
    display: flex;
    align-items: center;
    justify-content: space-between;
    padding: 0 40px;
    margin: 0 -40px 32px;
    width: calc(100% + 80px);
    position: sticky;
    top: 0;
    z-index: 1000;
    border-bottom: 3px solid var(--n-violet);
}

.n-topbar-left { display: flex; align-items: center; gap: 16px; }

.n-wordmark {
    font-family: 'Fraunces', serif;
    font-size: 21px;
    font-weight: 600;
    font-style: italic;
    color: #FFFFFF;
    letter-spacing: -0.2px;
    line-height: 1;
}

.n-divider-v { width: 1px; height: 20px; background: rgba(255,255,255,0.22); }

.n-app-name {
    font-size: 13px;
    font-weight: 500;
    color: rgba(255,255,255,0.82);
    letter-spacing: 0.01em;
}

.n-version {
    font-size: 11px;
    font-weight: 600;
    color: #FFFFFF;
    font-family: 'Inter', monospace;
    background: rgba(110,55,250,0.35);
    border: 1px solid rgba(110,55,250,0.6);
    padding: 4px 10px;
    border-radius: 20px;
}

/* ═══ BODY GRID ═══ */
.n-body {
    display: grid;
    grid-template-columns: 340px 1fr;
    min-height: calc(100vh - 60px);
    background: var(--n-paper);
}

/* ═══ LEFT PANEL ═══ */
.n-left {
    background: var(--n-panel);
    border: 1px solid var(--n-line);
    border-radius: 14px;
    padding: 32px 30px 36px;
    display: flex;
    flex-direction: column;
    box-shadow: 0 1px 2px rgba(32,28,44,0.03);
}

.n-section-label {
    font-size: 10.5px;
    font-weight: 700;
    letter-spacing: 0.13em;
    text-transform: uppercase;
    color: var(--n-ink-faint);
    margin-bottom: 20px;
    padding-bottom: 14px;
    border-bottom: 1px solid var(--n-line-soft);
}

.n-section-label::before {
    content: '●';
    color: var(--n-violet);
    font-size: 7px;
    margin-right: 7px;
    vertical-align: middle;
}

/* ═══ FILE BLOCK (uploader + one-line status, no duplicate label) ═══ */
.n-file-block { margin-bottom: 22px; }
.n-file-block:last-of-type { margin-bottom: 0; }

.n-file-status-line {
    display: flex;
    align-items: center;
    gap: 8px;
    margin-top: 8px;
    padding-left: 2px;
}

.n-file-dot {
    width: 7px; height: 7px;
    border-radius: 50%;
    background: #E4E1EC;
    flex-shrink: 0;
    transition: background 0.3s;
}

.n-file-dot.loaded { background: var(--n-good); box-shadow: 0 0 0 3px rgba(31,157,108,0.15); }
.n-file-dot.missing { background: #E4E1EC; }

.n-file-status { font-size: 11.5px; font-weight: 500; flex-shrink: 0; }
.n-file-status.loaded { color: var(--n-good); font-weight: 600; }
.n-file-status.missing { color: var(--n-ink-faint); }

/* ═══ PROGRESS ═══ */
.n-progress-block { margin: 28px 0 0; }

.n-progress-track {
    height: 5px;
    background: var(--n-line-soft);
    border-radius: 3px;
    overflow: hidden;
    margin-bottom: 10px;
}

.n-progress-fill {
    height: 100%;
    background: linear-gradient(90deg, var(--n-violet), var(--n-violet-dark));
    border-radius: 3px;
    transition: width 0.4s ease;
}

.n-progress-label { font-size: 11px; color: var(--n-ink-faint); font-weight: 500; }
.n-progress-label span { color: var(--n-ink); font-weight: 700; }

/* ═══ SPACER ═══ */
.n-spacer-24 { height: 24px; }
.n-spacer-16 { height: 16px; }
.n-spacer-12 { height: 12px; }
.n-spacer-8  { height: 8px; }

/* ═══ TOGGLE ROW (show logs) ═══ */
.n-toggle-row {
    display: flex;
    align-items: center;
    justify-content: flex-end;
    margin-bottom: -6px;
    padding-right: 2px;
}

.n-toggle-row [data-testid="stWidgetLabel"] p {
    font-size: 11.5px !important;
    font-weight: 600 !important;
    color: var(--n-ink-soft) !important;
    letter-spacing: 0.02em;
}

[data-testid="stToggle"] label div[data-baseweb="checkbox"] > div:first-child {
    background: #E4E1EC !important;
}

[data-testid="stToggle"] label[data-checked="true"] div[data-baseweb="checkbox"] > div:first-child,
[data-testid="stToggle"] div[aria-checked="true"] {
    background: var(--n-violet) !important;
}

/* ═══ RUN BUTTON OVERRIDE ═══ */
[data-testid="stButton"] > button {
    width: 100% !important;
    height: 48px !important;
    background: linear-gradient(135deg, var(--n-violet) 0%, var(--n-violet-dark) 100%) !important;
    color: #FFFFFF !important;
    font-family: 'Inter', sans-serif !important;
    font-size: 13px !important;
    font-weight: 700 !important;
    letter-spacing: 0.03em !important;
    border: none !important;
    border-radius: 9px !important;
    cursor: pointer !important;
    transition: transform 0.15s ease, box-shadow 0.2s !important;
    box-shadow: 0 3px 10px rgba(110,55,250,0.32), inset 0 1px 0 rgba(255,255,255,0.18) !important;
}

[data-testid="stButton"] > button:hover:not(:disabled) {
    transform: translateY(-1px);
    box-shadow: 0 6px 18px rgba(110,55,250,0.4), inset 0 1px 0 rgba(255,255,255,0.2) !important;
}

[data-testid="stButton"] > button:disabled {
    background: #F1EFF7 !important;
    color: #BBB6CC !important;
    box-shadow: none !important;
    cursor: not-allowed !important;
    border: 1px solid var(--n-line) !important;
}

/* ═══ FILE UPLOADER OVERRIDE ═══ */
[data-testid="stFileUploader"] { background: transparent !important; margin: 0 !important; padding: 0 !important; }

[data-testid="stFileUploader"] > div {
    background: #FCFBFA !important;
    border: 1.5px dashed #DAD5EA !important;
    border-radius: 8px !important;
    padding: 11px 14px !important;
    min-height: unset !important;
    transition: border-color 0.2s, background 0.2s !important;
}

[data-testid="stFileUploader"] > div:hover {
    border-color: var(--n-violet) !important;
    background: var(--n-violet-wash) !important;
}

[data-testid="stFileUploaderDropzoneInstructions"] {
    color: var(--n-ink-faint) !important;
    font-size: 11px !important;
    font-family: 'Inter', sans-serif !important;
}

[data-testid="stFileUploaderDropzoneInstructions"] > div > small { color: #C9C4DB !important; font-size: 10px !important; }

[data-testid="stFileUploaderFileName"] { color: var(--n-ink) !important; font-size: 11px !important; font-weight: 500 !important; }

[data-testid="stFileUploader"] label {
    font-size: 12px !important;
    font-weight: 600 !important;
    color: #55506A !important;
    font-family: 'Inter', sans-serif !important;
    margin-bottom: 6px !important;
}

/* ═══ RIGHT PANEL ═══ */
.n-right {
    background: var(--n-paper);
    padding: 32px 0 0;
    display: flex;
    flex-direction: column;
    gap: 28px;
}

/* ═══ PIPELINE BAR ═══ */
.n-pipeline-bar {
    background: var(--n-panel);
    border: 1px solid var(--n-line);
    border-radius: 14px;
    padding: 30px 34px;
    box-shadow: 0 1px 2px rgba(32,28,44,0.03);
}

.n-pipeline-bar-title {
    font-family: 'Fraunces', serif;
    font-size: 14px;
    font-style: italic;
    font-weight: 500;
    color: var(--n-ink-soft);
    margin-bottom: 26px;
}

.n-stages { display: grid; grid-template-columns: repeat(4, 1fr); gap: 16px; position: relative; }

.n-stages::before {
    content: '';
    position: absolute;
    top: 18px;
    left: calc(12.5% + 18px);
    right: calc(12.5% + 18px);
    height: 2px;
    background: var(--n-line-soft);
    z-index: 0;
}

.n-stage-col { display: flex; flex-direction: column; align-items: center; gap: 12px; position: relative; z-index: 1; }

.n-stage-icon {
    width: 36px; height: 36px;
    border-radius: 50%;
    border: 2px solid #E4E1EC;
    background: #FFFFFF;
    display: flex; align-items: center; justify-content: center;
    font-size: 14px;
    font-weight: 700;
    color: #C9C4DB;
    transition: all 0.3s ease;
}

.n-stage-col.done   .n-stage-icon { border-color: var(--n-good); background: var(--n-good); color: #FFFFFF; }
.n-stage-col.active .n-stage-icon { border-color: var(--n-violet); background: var(--n-violet); color: #FFFFFF; box-shadow: 0 0 0 4px rgba(110,55,250,0.14); }
.n-stage-col.failed .n-stage-icon { border-color: var(--n-fail); background: var(--n-fail-wash); color: var(--n-fail); }
.n-stage-col.skipped .n-stage-icon { border-color: var(--n-warn); background: var(--n-warn-wash); color: var(--n-warn); }

.n-stage-label { font-size: 12.5px; font-weight: 600; color: var(--n-ink-faint); text-align: center; transition: color 0.3s; margin-top: 2px; }
.n-stage-col.done   .n-stage-label { color: var(--n-good); }
.n-stage-col.active .n-stage-label { color: var(--n-violet); }
.n-stage-col.failed .n-stage-label { color: var(--n-fail); }
.n-stage-col.skipped .n-stage-label { color: var(--n-warn); }

.n-stage-sub { font-size: 10px; color: #C9C4DB; text-align: center; line-height: 1.4; }
.n-stage-col.done .n-stage-sub   { color: var(--n-ink-faint); }
.n-stage-col.active .n-stage-sub { color: var(--n-violet); opacity: 0.75; }

/* ═══ LOG CARD ═══ */
.n-log-card {
    background: var(--n-midnight);
    border-radius: 14px;
    overflow: hidden;
    flex: 1;
    box-shadow: 0 1px 2px rgba(32,28,44,0.03);
}

.n-log-header {
    padding: 18px 26px;
    border-bottom: 1px solid rgba(255,255,255,0.08);
    display: flex;
    align-items: center;
    justify-content: space-between;
    background: rgba(255,255,255,0.02);
}

.n-log-header-title {
    font-size: 11px;
    font-weight: 600;
    color: rgba(255,255,255,0.55);
    letter-spacing: 0.1em;
    text-transform: uppercase;
}

.n-log-dot { width: 7px; height: 7px; border-radius: 50%; background: rgba(255,255,255,0.18); }
.n-log-dot.active { background: var(--n-violet); box-shadow: 0 0 6px rgba(110,55,250,0.6); }
.n-log-dot.done   { background: var(--n-good); }

.n-log-body {
    padding: 24px 26px;
    font-family: 'SF Mono', 'Fira Code', 'Consolas', monospace;
    font-size: 12.5px;
    line-height: 2;
    min-height: 300px;
    max-height: 400px;
    overflow-y: auto;
}

.n-log-body::-webkit-scrollbar { width: 4px; }
.n-log-body::-webkit-scrollbar-track { background: rgba(255,255,255,0.02); }
.n-log-body::-webkit-scrollbar-thumb { background: rgba(255,255,255,0.15); border-radius: 2px; }

.ll { display: flex; gap: 12px; }
.ll-ts { color: rgba(255,255,255,0.22); flex-shrink: 0; }
.ll-ok   { color: #8FB8FF; }
.ll-good { color: #5CDDA0; }
.ll-fail { color: #FF7A8A; }
.ll-warn { color: #F0B84E; }
.ll-dim  { color: rgba(255,255,255,0.2); }
.ll-head { color: #B296FF; font-weight: 600; }
.ll-stage { color: rgba(255,255,255,0.35); }
.ll-div { border: none; border-top: 1px solid rgba(255,255,255,0.08); margin: 6px 0; }

/* ═══ COMPACT LOG (logs hidden) ═══ */
.n-log-compact {
    background: var(--n-panel);
    border: 1px solid var(--n-line);
    border-radius: 14px;
    padding: 24px 30px;
    display: flex;
    align-items: center;
    gap: 16px;
}

.n-log-compact-dot { width: 9px; height: 9px; border-radius: 50%; background: #E4E1EC; flex-shrink: 0; }
.n-log-compact-dot.active { background: var(--n-violet); box-shadow: 0 0 0 4px rgba(110,55,250,0.14); }
.n-log-compact-dot.done   { background: var(--n-good); box-shadow: 0 0 0 4px rgba(31,157,108,0.14); }
.n-log-compact-dot.fail   { background: var(--n-fail); box-shadow: 0 0 0 4px rgba(214,69,90,0.14); }

.n-log-compact-text { font-size: 13px; font-weight: 600; color: var(--n-ink); }
.n-log-compact-sub { font-size: 11px; color: var(--n-ink-faint); margin-top: 1px; }

/* ═══ EMPTY STATE ═══ */
.n-empty { display: flex; flex-direction: column; align-items: center; justify-content: center; gap: 10px; padding: 60px 40px; text-align: center; }
.n-empty-icon { font-size: 26px; opacity: 0.5; color: var(--n-violet); }
.n-empty-title { font-family: 'Fraunces', serif; font-size: 15px; font-weight: 500; font-style: italic; color: rgba(255,255,255,0.75); }
.n-empty-sub { font-size: 12px; color: rgba(255,255,255,0.35); max-width: 240px; line-height: 1.6; }

/* ═══ OUTPUT RESULTS ═══ */
.n-results { background: var(--n-panel); border: 1px solid var(--n-line); border-radius: 14px; padding: 30px 34px; box-shadow: 0 1px 2px rgba(32,28,44,0.03); }
.n-results-title { font-family: 'Fraunces', serif; font-size: 14px; font-style: italic; font-weight: 500; color: var(--n-ink-soft); margin-bottom: 24px; }
.n-result-grid { display: grid; grid-template-columns: repeat(4, 1fr); gap: 16px; margin-bottom: 26px; }

.n-result-card { border-radius: 10px; padding: 18px 20px; border: 1px solid var(--n-line-soft); background: #FCFBFA; }
.n-result-card.ok   { border-color: rgba(31,157,108,0.28); background: var(--n-good-wash); }
.n-result-card.fail { border-color: rgba(214,69,90,0.25);  background: var(--n-fail-wash); }
.n-result-card.skip { border-color: rgba(201,130,30,0.25); background: var(--n-warn-wash); }

.n-rc-name { font-size: 10px; font-weight: 700; letter-spacing: 0.1em; text-transform: uppercase; color: var(--n-ink-faint); margin-bottom: 10px; }
.n-rc-stat { font-size: 22px; font-weight: 800; line-height: 1; margin-bottom: 7px; }
.n-result-card.ok   .n-rc-stat { color: var(--n-good); }
.n-result-card.fail .n-rc-stat { color: var(--n-fail); }
.n-result-card.skip .n-rc-stat { color: var(--n-warn); }
.n-rc-sub { font-size: 10px; color: var(--n-ink-faint); line-height: 1.4; }

/* ═══ DOWNLOAD BUTTON ═══ */
[data-testid="stDownloadButton"] > button {
    height: 46px !important;
    background: var(--n-midnight) !important;
    color: #FFFFFF !important;
    font-family: 'Inter', sans-serif !important;
    font-size: 13px !important;
    font-weight: 600 !important;
    border: none !important;
    border-radius: 8px !important;
    padding: 0 24px !important;
    letter-spacing: 0.02em !important;
    cursor: pointer !important;
    transition: background 0.2s !important;
}

[data-testid="stDownloadButton"] > button:hover { background: var(--n-midnight-2) !important; }

/* Spinner */
[data-testid="stSpinner"] > div { color: var(--n-violet) !important; }

/* ═══ STREAMLIT FILE UPLOADER LABEL FIX ═══ */
.uploadedFileName { color: var(--n-ink) !important; font-size: 11px !important; }
</style>
""", unsafe_allow_html=True)

# ── Session state ─────────────────────────────────────────────────────────────
if "stage_states" not in st.session_state:
    st.session_state.stage_states = {1:"idle", 2:"idle", 3:"idle", 4:"idle"}
if "log_lines" not in st.session_state:
    st.session_state.log_lines = []
if "result_stats" not in st.session_state:
    st.session_state.result_stats = None
if "output_bytes" not in st.session_state:
    st.session_state.output_bytes = None
if "show_logs" not in st.session_state:
    st.session_state.show_logs = True

# ── Helpers ───────────────────────────────────────────────────────────────────
def ts():
    return time.strftime("%H:%M:%S")

def add_log(text, kind="ok"):
    st.session_state.log_lines.append((ts(), text, kind))

def set_stage(n, state):
    st.session_state.stage_states[n] = state

# ── Excel writer ──────────────────────────────────────────────────────────────
def write_output_styled(sheets: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    HDR_FILL  = PatternFill("solid", fgColor="1A1A1A")
    HDR_FONT  = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    HDR_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
    CELL_FONT = Font(name="Calibri", size=10)
    ODD_FILL  = PatternFill("solid", fgColor="F9F9F9")
    EVN_FILL  = PatternFill("solid", fgColor="FFFFFF")
    T_FILL    = PatternFill("solid", fgColor="F0FBF6"); T_FONT = Font(name="Calibri", size=10, color="22A06B", bold=True)
    F_fill    = PatternFill("solid", fgColor="FFF5F5"); F_FONT = Font(name="Calibri", size=10, color="CC0000", bold=True)
    thin      = Side(style="thin", color="EEEEEE")
    BORDER    = Border(bottom=thin)

    for name, df in sheets.items():
        ws = wb.create_sheet(title=name)
        ws.sheet_properties.tabColor = "FFFF00"

        for ci, col in enumerate(df.columns, 1):
            c = ws.cell(row=1, column=ci, value=str(col))
            c.fill=HDR_FILL; c.font=HDR_FONT; c.alignment=HDR_ALIGN; c.border=BORDER

        for ri, row in enumerate(df.itertuples(index=False), 2):
            fill = ODD_FILL if ri % 2 == 0 else EVN_FILL
            for ci, val in enumerate(row, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                cname = df.columns[ci-1]
                if cname == "Check":
                    if val is True:
                        c.fill=T_FILL; c.font=T_FONT; c.value="TRUE"
                    elif val is False:
                        c.fill=F_fill; c.font=F_FONT; c.value="FALSE"
                    else:
                        c.fill=fill; c.font=CELL_FONT
                else:
                    c.fill=fill; c.font=CELL_FONT

        for col in ws.columns:
            ltr = get_column_letter(col[0].column)
            mx = max((len(str(c.value)) for c in col if c.value), default=8)
            ws.column_dimensions[ltr].width = min(max(mx + 2, 8), 36)

        ws.freeze_panes = "A2"
        if df.shape[1] > 0:
            ws.auto_filter.ref = f"A1:{get_column_letter(df.shape[1])}1"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

# ── Pipeline runner ───────────────────────────────────────────────────────────
def run_pipeline(adapt_f, bsr_f, yt_f, matex_f, sample_f):
    st.session_state.log_lines = []
    st.session_state.result_stats = None
    st.session_state.output_bytes = None
    for n in [1,2,3,4]: set_stage(n, "idle")

    add_log("QC CHECK AUTOMATION PIPELINE", "head")
    add_log(f"Started at {time.strftime('%Y-%m-%d %H:%M:%S')}", "dim")
    add_log("", "dim")
    add_log("── LOAD ──────────────────────────────────────────────", "stage")

    load_ok = True
    frames  = {}

    for key, uf, sheet, hdr, label in [
        ("adapt",  adapt_f,  SHEETS["adapt"],  0, "Adapt"),
        ("bsr",    bsr_f,    SHEETS["bsr"],    5, "BSR"),
        ("yt",     yt_f,     SHEETS["yt"],     0, "YouTube MM"),
        ("matex",  matex_f,  SHEETS["matex"],  0, "Matex"),
        ("sample", sample_f, SHEETS["sample"], 0, "Sample"),
    ]:
        try:
            buf = io.BytesIO(uf.read())
            df  = pd.read_excel(buf, sheet_name=sheet, header=hdr, engine="openpyxl")
            df.columns = [str(c).strip() for c in df.columns]
            frames[key] = df
            add_log(f"  [✓]  {label:<14}  {len(df):,} rows · {len(df.columns)} cols", "good")
        except Exception as e:
            add_log(f"  [✗]  {label:<14}  {e}", "fail")
            load_ok = False

    if not load_ok:
        add_log("", "dim"); add_log("  PIPELINE ABORTED — fix load errors above", "fail")
        return

    add_log("", "dim")
    add_log("── STAGES ────────────────────────────────────────────", "stage")

    output_sheets  = {}
    sample_summary = None
    stats = {}

    # Stage 1
    set_stage(1, "active")
    try:
        df = build_rate_check(frames["adapt"], frames["bsr"], frames["yt"])
        output_sheets["RateCheck"] = df
        stats[1] = {"rows": len(df), "ok": True}
        set_stage(1, "done")
        add_log(f"  [✓]  RateCheck        {len(df):,} rows", "good")
    except Exception as e:
        set_stage(1, "failed"); stats[1] = {"ok": False, "error": str(e)}
        add_log(f"  [✗]  RateCheck        {e}", "fail")

    # Stage 2
    set_stage(2, "active")
    try:
        df = build_exposure_check(frames["sample"], frames["matex"])
        output_sheets["ExposureCheck"] = df
        stats[2] = {"rows": len(df), "ok": True}
        set_stage(2, "done")
        add_log(f"  [✓]  ExposureCheck    {len(df):,} rows", "good")
    except Exception as e:
        set_stage(2, "failed"); stats[2] = {"ok": False, "error": str(e)}
        add_log(f"  [✗]  ExposureCheck    {e}", "fail")

    # Stage 3
    set_stage(3, "active")
    try:
        ss, sample_summary = build_sample(frames["sample"])
        output_sheets["Sample"] = ss
        stats[3] = {"rows": len(ss), "matchdays": len(sample_summary), "ok": True}
        set_stage(3, "done")
        add_log(f"  [✓]  Sample           {len(ss):,} rows · {len(sample_summary)} matchdays", "good")
    except Exception as e:
        set_stage(3, "failed"); stats[3] = {"ok": False, "error": str(e)}
        add_log(f"  [✗]  Sample           {e}", "fail")

    # Stage 4
    set_stage(4, "active")
    try:
        if sample_summary is None:
            raise RuntimeError("Stage 3 must succeed before Stage 4 can run")
        df = build_extrap_check(frames["adapt"], sample_summary)
        output_sheets["ExtrapCheck"] = df
        stats[4] = {"rows": len(df), "ok": True}
        set_stage(4, "done")
        add_log(f"  [✓]  ExtrapCheck      {len(df):,} rows", "good")
    except RuntimeError as e:
        set_stage(4, "skipped"); stats[4] = {"ok": False, "skipped": True, "error": str(e)}
        add_log(f"  [⚠]  ExtrapCheck      SKIPPED — {e}", "warn")
    except Exception as e:
        set_stage(4, "failed"); stats[4] = {"ok": False, "error": str(e)}
        add_log(f"  [✗]  ExtrapCheck      {e}", "fail")

    add_log("", "dim")
    add_log("── OUTPUT ────────────────────────────────────────────", "stage")
    done = sum(1 for s in stats.values() if s.get("ok"))

    if output_sheets:
        try:
            st.session_state.output_bytes = write_output_styled(output_sheets)
            add_log(f"  [✓]  {done}/4 sheets written → QC_Output.xlsx", "good")
        except Exception as e:
            add_log(f"  [✗]  Write failed: {e}", "fail")

    add_log("", "dim")
    if done == 4:
        add_log("  ✓  COMPLETE — all 4 sheets written", "good")
    elif done > 0:
        add_log(f"  ⚠  PARTIAL — {done}/4 sheets written", "warn")
    else:
        add_log("  ✗  FAILED — no output written", "fail")

    st.session_state.result_stats = stats


# ══════════════════════════════════════════════════════════════════════════════
# RENDER
# ══════════════════════════════════════════════════════════════════════════════

# ── TOPBAR ───────────────────────────────────────────────────────────────────
st.markdown("""
<div class="n-topbar">
  <div class="n-topbar-left">
    <span class="n-wordmark">Nielsen</span>
    <div class="n-divider-v"></div>
    <span class="n-app-name">QC Automation Pipeline</span>
  </div>
  <span class="n-version">v7.0 · Media Measurement</span>
</div>
""", unsafe_allow_html=True)

# ── COLUMNS ──────────────────────────────────────────────────────────────────
left_col, right_col = st.columns([0.9, 2.1], gap="small")

# ═══════════════
# LEFT
# ═══════════════
with left_col:
    st.markdown('<div class="n-left">', unsafe_allow_html=True)

    # ── Files ────────────────────────────────────────────────────────────────
    st.markdown('<div class="n-section-label">Source Files</div>', unsafe_allow_html=True)

    FILE_DEFS = [
        ("adapt",  "Adapt",      "programs"),
        ("bsr",    "BSR",        "Database · row 6"),
        ("yt",     "YouTube MM", "MM"),
        ("matex",  "Matex",      "exposures"),
        ("sample", "Sample",     "programs"),
    ]

    uploaded = {}
    loaded_count = 0

    for key, label, hint in FILE_DEFS:
        st.markdown('<div class="n-file-block">', unsafe_allow_html=True)
        uf = st.file_uploader(
            f"{label}",
            type=["xlsx"],
            key=f"upload_{key}",
            help=f"Sheet: {hint}",
            label_visibility="visible",
        )
        uploaded[key] = uf
        dot_cls = "loaded" if uf else "missing"
        status_cls = "loaded" if uf else "missing"
        status_txt = f"✓  {uf.size/1024:.0f} KB · sheet: {hint}" if uf else f"Awaiting upload · sheet: {hint}"
        if uf:
            loaded_count += 1
        st.markdown(f"""
        <div class="n-file-status-line">
          <div class="n-file-dot {dot_cls}"></div>
          <div class="n-file-status {status_cls}">{status_txt}</div>
        </div>
        </div>
        """, unsafe_allow_html=True)

    # Progress
    pct = int(loaded_count / 5 * 100)
    st.markdown(f"""
    <div class="n-progress-block">
      <div class="n-progress-track">
        <div class="n-progress-fill" style="width:{pct}%"></div>
      </div>
      <div class="n-progress-label"><span>{loaded_count}</span> of 5 files ready</div>
    </div>
    <div class="n-spacer-24"></div>
    """, unsafe_allow_html=True)

    # Run button
    all_ready = loaded_count == 5
    btn_label = "Run QC Pipeline →" if all_ready else f"Waiting for {5 - loaded_count} more file{'s' if 5-loaded_count != 1 else ''}"
    if st.button(btn_label, disabled=not all_ready, use_container_width=True):
        with st.spinner("Running…"):
            run_pipeline(
                uploaded["adapt"], uploaded["bsr"], uploaded["yt"],
                uploaded["matex"], uploaded["sample"]
            )
        st.rerun()

    st.markdown('</div>', unsafe_allow_html=True)


# ═══════════════
# RIGHT
# ═══════════════
with right_col:
    st.markdown('<div class="n-right">', unsafe_allow_html=True)

    # ── Pipeline bar ─────────────────────────────────────────────────────────
    STAGE_DEFS = [
        (1, "RateCheck",     "Adapt · BSR · YT"),
        (2, "ExposureCheck", "Sample vs Matex"),
        (3, "Sample",        "pID · matchdays"),
        (4, "ExtrapCheck",   "BT% / Msec%"),
    ]

    ICON_MAP = {"idle":"—","done":"✓","active":"◎","failed":"✗","skipped":"⚠"}

    cols_html = ""
    for n, title, sub in STAGE_DEFS:
        state = st.session_state.stage_states[n]
        icon  = ICON_MAP.get(state, "—")

        # Enrich sub with row count after run
        if state == "done" and st.session_state.result_stats:
            s = st.session_state.result_stats.get(n, {})
            rows = s.get("rows")
            if rows:
                sub = f"{rows:,} rows"
            if n == 3 and s.get("matchdays"):
                sub += f" · {s['matchdays']} matchdays"

        cols_html += f"""
        <div class="n-stage-col {state}">
          <div class="n-stage-icon">{icon}</div>
          <div class="n-stage-label">{title}</div>
          <div class="n-stage-sub">{sub}</div>
        </div>"""

    st.markdown(f"""
    <div class="n-pipeline-bar">
      <div class="n-pipeline-bar-title">Pipeline Stages</div>
      <div class="n-stages">{cols_html}</div>
    </div>
    """, unsafe_allow_html=True)

    # ── Log visibility toggle ────────────────────────────────────────────────
    done_count = sum(1 for s in (st.session_state.result_stats or {}).values() if s.get("ok"))
    dot_class = "active" if st.session_state.log_lines and not st.session_state.result_stats else ("done" if st.session_state.result_stats else "")

    st.markdown('<div class="n-toggle-row">', unsafe_allow_html=True)
    st.session_state.show_logs = st.toggle(
        "Show pipeline logs",
        value=st.session_state.show_logs,
        key="show_logs_toggle",
    )
    st.markdown('</div>', unsafe_allow_html=True)

    # ── Log ──────────────────────────────────────────────────────────────────
    if st.session_state.show_logs:
        if not st.session_state.log_lines:
            log_content = """
            <div class="n-empty">
              <div class="n-empty-icon">○</div>
              <div class="n-empty-title">No pipeline run yet</div>
              <div class="n-empty-sub">Upload all 5 source files and press Run QC Pipeline</div>
            </div>"""
        else:
            rows_html = ""
            for stamp, text, kind in st.session_state.log_lines:
                if text == "":
                    rows_html += '<div style="height:3px"></div>'
                elif text.startswith("──"):
                    rows_html += f'<hr class="ll-div"><div class="ll"><span class="ll-ts">{stamp}</span><span class="ll-stage">{text}</span></div>'
                else:
                    cls = {"ok":"ll-ok","good":"ll-good","fail":"ll-fail","warn":"ll-warn","dim":"ll-dim","head":"ll-head","stage":"ll-stage"}.get(kind,"ll-ok")
                    rows_html += f'<div class="ll"><span class="ll-ts">{stamp}</span><span class="{cls}">{text}</span></div>'
            log_content = f'<div class="n-log-body">{rows_html}</div>'

        st.markdown(f"""
        <div class="n-log-card">
          <div class="n-log-header">
            <span class="n-log-header-title">Pipeline Log</span>
            <div class="n-log-dot {dot_class}"></div>
          </div>
          {log_content}
        </div>
        """, unsafe_allow_html=True)
    else:
        # Compact status strip when logs are hidden
        if not st.session_state.log_lines:
            compact_dot, compact_text, compact_sub = "", "No pipeline run yet", "Toggle logs on for step-by-step detail"
        elif st.session_state.result_stats and done_count == 4:
            compact_dot, compact_text, compact_sub = "done", "Complete — all 4 sheets written", "Toggle logs on for step-by-step detail"
        elif st.session_state.result_stats and done_count > 0:
            compact_dot, compact_text, compact_sub = "fail", f"Partial — {done_count}/4 sheets written", "Toggle logs on to see what failed"
        elif st.session_state.result_stats:
            compact_dot, compact_text, compact_sub = "fail", "Failed — no output written", "Toggle logs on to see what failed"
        else:
            compact_dot, compact_text, compact_sub = "active", "Running…", "Toggle logs on for step-by-step detail"

        st.markdown(f"""
        <div class="n-log-compact">
          <div class="n-log-compact-dot {compact_dot}"></div>
          <div>
            <div class="n-log-compact-text">{compact_text}</div>
            <div class="n-log-compact-sub">{compact_sub}</div>
          </div>
        </div>
        """, unsafe_allow_html=True)

    # ── Results + Download ────────────────────────────────────────────────────
    if st.session_state.result_stats:
        stats = st.session_state.result_stats
        names = ["RateCheck","ExposureCheck","Sample","ExtrapCheck"]

        cards_html = ""
        for i, sname in enumerate(names, 1):
            s = stats.get(i, {})
            ok      = s.get("ok", False)
            skipped = s.get("skipped", False)
            cls     = "ok" if ok else ("skip" if skipped else "fail")
            icon    = "✓" if ok else ("⚠" if skipped else "✗")
            rows    = s.get("rows")
            if ok:
                sub = f"{rows:,} rows" if rows else "complete"
                if i == 3 and s.get("matchdays"): sub += f"<br>{s['matchdays']} matchdays"
            elif skipped:
                sub = "Stage 3 required"
            else:
                sub = "See log above"

            cards_html += f"""
            <div class="n-result-card {cls}">
              <div class="n-rc-name">{sname}</div>
              <div class="n-rc-stat">{icon}</div>
              <div class="n-rc-sub">{sub}</div>
            </div>"""

        st.markdown(f"""
        <div class="n-results">
          <div class="n-results-title">Output Sheets</div>
          <div class="n-result-grid">{cards_html}</div>
        </div>
        """, unsafe_allow_html=True)

        if st.session_state.output_bytes:
            fname = f"QC_Output_{time.strftime('%Y%m%d_%H%M')}.xlsx"
            st.download_button(
                label=f"↓  Download {fname}",
                data=st.session_state.output_bytes,
                file_name=fname,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

    st.markdown('</div>', unsafe_allow_html=True)
